# -*- coding: utf-8 -*-
"""
分镜导出器（确定性）：storyboard.json -> 分镜.md / 分镜.xlsx / 提示词.md
用法: python export_docs.py <storyboard.json> [输出目录]
"""
import sys,json,os
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter

MODE_CN={"follow":"跟随跟拍","cu":"特写推镜(固定)","ots":"越肩OTS(固定)","wide":"广角双人","keys":"关键帧运镜"}

def load(p): return json.load(open(p,encoding="utf-8"))

def shot_rows(cfg):
    rows=[]; t=0.0
    for sh in cfg["shots"]:
        rows.append((sh,t,t+sh["dur"])); t+=sh["dur"]
    return rows

def to_md(cfg):
    L=[]
    L.append(f"# {cfg['title']}\n")
    L.append(f"> 项目 `{cfg['project']}` ｜ {cfg.get('w',960)}×{cfg.get('h',540)} ｜ {cfg.get('fps',24)}fps ｜ 场景 `{cfg['scene']}` ｜ 共 {len(cfg['shots'])} 镜 / {sum(s['dur'] for s in cfg['shots']):.1f}s\n")
    L.append("## 人物\n")
    L.append("| ID | 角色 | 服装色RGB | 围裙 | 站位/走位 |")
    L.append("|---|---|---|---|---|")
    for aid,a in cfg["actors"].items():
        pos = "走位 path（主角）" if aid=="c" else f"固定 {a.get('pos')}"
        L.append(f"| {aid} | {a.get('name','')} | {a.get('shirt')} | {a.get('apron')} | {pos} |")
    L.append("\n## 分镜表\n")
    L.append("| 镜号 | 时间(s) | 景别 | 运镜 | 台词/字幕 | 中文提示词 |")
    L.append("|---|---|---|---|---|---|")
    for sh,a,b in shot_rows(cfg):
        L.append(f"| {sh['id']} | {a:.1f}-{b:.1f} | {sh.get('shot','')} | {MODE_CN.get(sh['mode'],sh['mode'])} | {sh.get('line','')} | {sh.get('prompt_cn','')} |")
    if any(s.get("prompt_en") for s in cfg["shots"]):
        L.append("\n## English Prompts\n")
        for sh,a,b in shot_rows(cfg):
            if sh.get("prompt_en"): L.append(f"- **{sh['id']}** ({a:.1f}-{b:.1f}s, {MODE_CN.get(sh['mode'],sh['mode'])}): {sh['prompt_en']}")
    return "\n".join(L)+"\n"

def to_prompts(cfg):
    L=[f"# {cfg['title']} · 分镜提示词（Seedance/可灵）\n"]
    for sh,a,b in shot_rows(cfg):
        L.append(f"### {sh['id']}（{a:.1f}-{b:.1f}s · {sh.get('shot','')} · {MODE_CN.get(sh['mode'],sh['mode'])}）")
        if sh.get("line"): L.append(f"- 台词/字幕：{sh['line']}")
        if sh.get("prompt_cn"): L.append(f"- 🇨🇳 {sh['prompt_cn']}")
        if sh.get("prompt_en"): L.append(f"- 🇬🇧 `{sh['prompt_en']}`")
        L.append("")
    return "\n".join(L)

def to_xlsx(cfg,path):
    wb=Workbook(); ws=wb.active; ws.title="分镜脚本"
    hdr=PatternFill("solid",fgColor="2F5233"); hf=Font(color="FFFFFF",bold=True,name="微软雅黑")
    cf=Font(size=10,name="微软雅黑"); thin=Side(style="thin",color="BBBBBB")
    bd=Border(left=thin,right=thin,top=thin,bottom=thin); wrap=Alignment(wrap_text=True,vertical="top")
    cols=["镜号","开始(s)","结束(s)","景别","运镜","台词/字幕","中文提示词","English"]
    ws.append(cols)
    for sh,a,b in shot_rows(cfg):
        ws.append([sh["id"],round(a,1),round(b,1),sh.get("shot",""),MODE_CN.get(sh["mode"],sh["mode"]),
                   sh.get("line",""),sh.get("prompt_cn",""),sh.get("prompt_en","")])
    for c in range(1,len(cols)+1):
        cell=ws.cell(1,c); cell.fill=hdr; cell.font=hf; cell.alignment=Alignment(horizontal="center",vertical="center"); cell.border=bd
    for i,w in enumerate([16,9,9,10,16,30,46,46],1): ws.column_dimensions[get_column_letter(i)].width=w
    for r in range(2,ws.max_row+1):
        for c in range(1,len(cols)+1):
            cell=ws.cell(r,c); cell.font=cf; cell.border=bd; cell.alignment=wrap
    ws.freeze_panes="A2"
    wb.save(path)

if __name__=="__main__":
    cfg=load(sys.argv[1])
    outdir=sys.argv[2] if len(sys.argv)>2 else os.path.join(os.path.dirname(sys.argv[1]),"..","out")
    outdir=os.path.abspath(outdir); os.makedirs(outdir,exist_ok=True)
    pj=cfg["project"]
    md=os.path.join(outdir,pj+"_分镜脚本.md"); open(md,"w",encoding="utf-8").write(to_md(cfg))
    pp=os.path.join(outdir,pj+"_提示词.md"); open(pp,"w",encoding="utf-8").write(to_prompts(cfg))
    xls=os.path.join(outdir,pj+"_分镜脚本.xlsx"); to_xlsx(cfg,xls)
    print(f"[OK] 导出:\n  {md}\n  {pp}\n  {xls}")


