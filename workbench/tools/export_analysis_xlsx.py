# -*- coding: utf-8 -*-
"""analysis.json -> 分镜脚本 Excel(.xlsx)（AI 短片分析工作台 阶段三）
用法: python export_analysis_xlsx.py <analysis.json> [out.xlsx]
sheet1「分镜脚本」: 镜号/时间码/时长/景别/运镜/角度/光线/剧情/动作/台词/AI提示词
  （E08：跨镜台词在非主属镜的引用带 span=overlap，台词 cell 加 ↔ 前缀）
sheet2「台词汇总」: 时间/角色/文本（有顶层 dialogue_track 时按本体展开——完整时间、
  一句一行天然去重；无 track 按镜头顺序展开并按 event 去重）
默认输出: <analysis目录>/<name>_分镜脚本.xlsx"""
import sys, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

def fmt_tc(t):
    """秒 -> mm:ss.s 时间码"""
    m = int(t // 60); s = t - m * 60
    return f"{m:02d}:{s:04.1f}"

def main():
    src = sys.argv[1]
    a = json.load(open(src, encoding="utf-8"))
    out = sys.argv[2] if len(sys.argv) > 2 else \
        os.path.join(os.path.dirname(src), (a.get("name") or os.path.splitext(os.path.basename(src))[0]) + "_分镜脚本.xlsx")
    shots = a.get("shots") or []

    wb = Workbook(); ws = wb.active; ws.title = "分镜脚本"
    head = ["镜号", "时间码", "时长(s)", "景别", "运镜", "角度", "光线", "剧情", "动作", "台词", "AI提示词"]
    hf = Font(bold=True, color="FFFFFF", size=11, name="微软雅黑"); hfill = PatternFill("solid", fgColor="2F5233")
    thin = Side(style="thin", color="BBBBBB"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    ws.append(head)
    for i in range(1, len(head) + 1):
        c = ws.cell(1, i); c.font = hf; c.fill = hfill
        c.alignment = Alignment(horizontal="center", vertical="center"); c.border = border
    line_rows = []
    for idx, s in enumerate(shots, 1):
        t0, t1 = s.get("t_in", 0), s.get("t_out", 0)
        dl = s.get("dialogue") or []
        # E08：span=overlap 是跨镜台词的非主属镜引用，加 ↔ 前缀提示"这句在别镜也在"
        line_txt = "\n".join(("↔" if d.get("span") == "overlap" else "")
                             + f"【{d.get('speaker','?')}】{d.get('text','')}" for d in dl)
        row = [s.get("id", f"S{idx}"), f"{fmt_tc(t0)}–{fmt_tc(t1)}", round(s.get("duration", t1 - t0), 2),
               s.get("shot_size", ""), s.get("camera_move", ""), s.get("angle", ""),
               s.get("lighting", ""), s.get("story", ""), s.get("action", ""), line_txt,
               s.get("prompt_cn", "")]
        ws.append(row)
        for c in range(1, len(head) + 1):
            cell = ws.cell(idx + 1, c); cell.alignment = wrap; cell.border = border
            cell.font = Font(size=10, name="微软雅黑")
        for d in dl:
            line_rows.append([fmt_tc(d.get("t_in", t0)) + "–" + fmt_tc(d.get("t_out", t1)),
                              d.get("speaker", "?"), d.get("text", ""), d.get("event")])
    for i, w in enumerate([7, 15, 8, 8, 8, 8, 20, 34, 28, 34, 46], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"

    # sheet2 台词汇总：有 dialogue_track（E08 独立音轨本体）时按本体展开——完整时间、
    # 一句一行天然去重（跨镜句不再因挂多镜而重复）；无 track 走原 per-shot 展开，
    # 并按 event 去重防御（同一 event 在相邻镜重复出现时只保留首次）。
    track = a.get("dialogue_track")
    if isinstance(track, list) and track:
        line_rows = [[fmt_tc(e.get("t_in", 0)) + "–" + fmt_tc(e.get("t_out", 0)),
                      e.get("speaker", "?"), e.get("text", ""), e.get("event")] for e in track]
    else:
        seen = set(); dedup = []
        for r in line_rows:
            ev = r[3]
            if ev:
                if ev in seen:
                    continue
                seen.add(ev)
            dedup.append(r)
        line_rows = dedup

    ws2 = wb.create_sheet("台词汇总")
    h2 = ["时间", "角色", "文本"]
    ws2.append(h2)
    for i in range(1, 4):
        c = ws2.cell(1, i); c.font = hf; c.fill = hfill
        c.alignment = Alignment(horizontal="center", vertical="center"); c.border = border
    for r in line_rows:
        ws2.append(r[:3])   # 第 4 列 event 仅内部去重用，不写进 Excel
        for c in range(1, 4):
            cell = ws2.cell(ws2.max_row, c); cell.alignment = wrap; cell.border = border
            cell.font = Font(size=10, name="微软雅黑")
    for col, w in zip("ABC", [16, 10, 60]):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"

    wb.save(out)
    print(f"saved: {out}（分镜 {len(shots)} 行，台词 {len(line_rows)} 行）")
    print(f"OUTPUT:{out}")

if __name__ == "__main__":
    main()
