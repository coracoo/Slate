# -*- coding: utf-8 -*-
"""③ 网页表格与 Excel 导出的"名字"口径对齐回归（业务链路审计 §六.13 / 待你定 12）。

用户 09-25 定版：**网页补景别列 + 两侧都用名字**。原状三处不一致：
* 网页没有景别列（xlsx 有）；
* xlsx「场景」列直接吐 `room/field` 原值（网页显示场景名）；
* 网页台词列显示说话人 **id**（xlsx 显示名字）。
本文件锁住 xlsx 侧的解析链：场景=scene_ref→场景档案名，没关联才回落 室内/外景；
说话人=板内 actors → ② 人物档案 → 才退 id。前端同一条链在 `utils/assetNames.ts`。
"""
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
TOOL = ROOT / 'workbench' / 'tools' / 'export_storyboard_xlsx.py'


class XlsxNameColumnsTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.proj = Path(self._tmp.name)
        (self.proj / '分镜').mkdir(parents=True)
        (self.proj / '素材').mkdir(parents=True)
        self.write('素材/场景.json', {"scenes": [
            {"id": "loc_tent", "name": "中军帐", "aliases": ["军帐"]},
            {"id": "loc_field", "name": "野外战场"}]})
        self.write('素材/人物.json', {"characters": [
            {"id": "hero", "name": "主角"}, {"id": "lu", "name": "刘备"}]})
        self.board = self.write('分镜/剧本_E1.json', {"title": "E1", "actors": {"hero": {"name": "板内名·主角"}}, "shots": [
            {"id": "S1", "scene_ref": "@scene:loc_tent", "scene": "room", "dur": 4, "shot_size": "中景",
             "lines": [{"at": 0, "speaker": "hero", "line": "第一句"}, {"at": 2, "speaker": "lu", "line": "第二句"}]},
            {"id": "S2", "scene": "field", "dur": 3, "shot_size": "全景",
             "lines": [{"at": 0, "speaker": "narrator", "line": "旁白一句"}]},
            {"id": "S3", "scene": "room", "dur": 2, "lines": []},
        ]})

    def tearDown(self):
        self._tmp.cleanup()

    def write(self, rel, payload):
        path = self.proj / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
        return str(path)

    def _rows(self):
        out = str(self.proj / 'out.xlsx')
        env = dict(os.environ, PYTHONPATH=str(ROOT / 'workbench' / 'tools'))
        done = subprocess.run([sys.executable, str(TOOL), self.board, out],
                              stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                              text=True, encoding='utf-8', errors='replace', env=env)
        self.assertEqual(done.returncode, 0, done.stdout)
        from openpyxl import load_workbook
        wb = load_workbook(out)
        sheet = wb['分镜脚本']
        header = [c.value for c in sheet[1]]
        return header, [[c.value for c in row] for row in sheet.iter_rows(min_row=2)], \
            [[c.value for c in row] for row in wb['台词汇总'].iter_rows(min_row=2)]

    def test_scene_column_uses_names_not_room_field(self):
        header, rows, _ = self._rows()
        scene = header.index('场景')
        self.assertEqual([r[scene] for r in rows], ['中军帐', '外景', '室内'],
                         '场景列必须与 ③ 网页一致：scene_ref 出场景名，没关联才回落 室内/外景')
        self.assertNotIn('room', [r[scene] for r in rows])
        self.assertNotIn('field', [r[scene] for r in rows])

    def test_speaker_prefers_board_then_archive_then_id(self):
        header, rows, _ = self._rows()
        dlg = rows[0][header.index('台词')]
        self.assertIn('【板内名·主角】', dlg, '板内 actors 是这一本分镜自带的名字表，优先级最高')
        self.assertIn('【刘备】', dlg, '板内没有的说话人要能从 ② 人物档案解名，不能吐 id')
        _, _, sheet2 = self._rows()
        self.assertEqual([r[2] for r in sheet2], ['板内名·主角', '刘备', 'narrator'],
                         '台词汇总 sheet 同一条链；旁白按契约不建档，解不到名就照实留 narrator')

    def test_shot_size_column_present_and_aligned_with_web(self):
        header, rows, _ = self._rows()
        self.assertIn('景别', header)
        self.assertEqual(header.index('景别'), header.index('场景') + 1,
                         '列序与 ③ 网页一致（镜号/场景/景别/时长s…），交接时两双手不用来回找')
        self.assertEqual([r[header.index('景别')] for r in rows], ['中景', '全景', None])

    def test_missing_archive_does_not_break_export(self):
        """没提炼过素材的老项目：场景/人物档案都不存在，导出照旧出表，不报错也不吐 Python 字典。"""
        os.remove(self.proj / '素材' / '场景.json')
        os.remove(self.proj / '素材' / '人物.json')
        header, rows, _ = self._rows()
        scene, dlg = header.index('场景'), header.index('台词')
        self.assertEqual(rows[0][scene], 'loc_tent', '档案没了就照实显示 ref，不编名字')
        self.assertIn('【板内名·主角】', rows[0][dlg], '板内 actors 是分镜自带的，不依赖 ② 档案')
        self.assertIn('【lu】', rows[0][dlg], '档案删了才解不到名——这时照实退回 id，不静默丢台词')


if __name__ == '__main__':
    unittest.main()
