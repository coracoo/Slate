# -*- coding: utf-8 -*-
"""拉片证据层回归测试（update.md E02 切点溯源 / E03 「不确定」合法出口 + E03 增强转场交界帧）。

E02: merge_short_shots 返回合并映射（merged_from 原始边界）；detect_shots 返回检测方式；
     build_cut_detection 对均匀 5s 兜底标 sampling:true。
E03: 受控词表含合法值「不确定」；validate_analysis 接受它且对 sampling 醒目警告（不阻断）；
     analysis_to_storyboard 对「不确定」/空值落默认几何（中景/固定/平视）不崩。
E03 增强（转场交界帧）: 每镜（除首镜）加抽 S{n}_prev_tail/S{n}_head 交界帧，AI 依据
     「上一镜结尾 vs 本镜开头」证据判断 transition；帧缺失时退化允许「不确定」。
"""
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "workbench"))
sys.path.insert(0, str(ROOT / "workbench" / "tools"))
sys.path.insert(0, str(ROOT / "previs_system" / "tools"))

import analyze_film
import analysis_to_storyboard as a2s
import validate_analysis


class MergeShortShotsTests(unittest.TestCase):
    """碎镜合并映射：merged_from 只记录实际吞并过邻居的输出段。"""

    def test_min_dur_zero_no_overhead(self):
        segs = [(0.0, 0.5), (0.5, 1.0)]
        out, mf = analyze_film.merge_short_shots(segs, 0)
        self.assertEqual(out, segs)      # 不合并
        self.assertEqual(mf, {})         # 零开销：无 merged_from

    def test_min_dur_none_no_overhead(self):
        segs = [(0.0, 0.5), (0.5, 1.0)]
        out, mf = analyze_film.merge_short_shots(segs, None)
        self.assertEqual(out, segs)
        self.assertEqual(mf, {})

    def test_all_long_no_merge(self):
        segs = [(0.0, 3.0), (3.0, 6.0)]
        out, mf = analyze_film.merge_short_shots(segs, 2.0)
        self.assertEqual(out, segs)
        self.assertEqual(mf, {})         # 未合并的段不出现

    def test_head_segment_merges_into_only_neighbor(self):
        out, mf = analyze_film.merge_short_shots([(0.0, 0.8), (0.8, 6.0)], 2.0)
        self.assertEqual(out, [(0.0, 6.0)])
        self.assertEqual(mf, {0: [[0.0, 0.8], [0.8, 6.0]]})

    def test_tail_segment_merges_into_only_neighbor(self):
        out, mf = analyze_film.merge_short_shots([(0.0, 5.0), (5.0, 5.8)], 2.0)
        self.assertEqual(out, [(0.0, 5.8)])
        self.assertEqual(mf, {0: [[0.0, 5.0], [5.0, 5.8]]})

    def test_four_short_shots_keep_all_original_bounds(self):
        # 4×0.8s 碎镜夹在两个长镜之间：全部并入左邻，原始切点边界完整保留
        segs = [(0.0, 5.0), (5.0, 5.8), (5.8, 6.6), (6.6, 7.4), (7.4, 8.2), (8.2, 15.0)]
        out, mf = analyze_film.merge_short_shots(segs, 2.0)
        self.assertEqual(out, [(0.0, 8.2), (8.2, 15.0)])
        self.assertEqual(mf, {0: [[0.0, 5.0], [5.0, 5.8], [5.8, 6.6], [6.6, 7.4], [7.4, 8.2]]})

    def test_merge_into_longer_right_neighbor_with_recheck(self):
        # 右邻更长 -> 并入右邻，并回退重查被撑大的邻居
        out, mf = analyze_film.merge_short_shots([(0.0, 3.0), (3.0, 3.8), (3.8, 10.0)], 2.0)
        self.assertEqual(out, [(0.0, 3.0), (3.0, 10.0)])
        self.assertEqual(mf, {1: [[3.0, 3.8], [3.8, 10.0]]})

    def test_head_chain_collapses_until_long_enough(self):
        # 首段连续吞并后仍短于阈值，继续向后并（i==0 链式路径）
        out, mf = analyze_film.merge_short_shots([(0.0, 1.0), (1.0, 1.5), (1.5, 9.0)], 2.0)
        self.assertEqual(out, [(0.0, 9.0)])
        self.assertEqual(mf, {0: [[0.0, 1.0], [1.0, 1.5], [1.5, 9.0]]})


class CutDetectionTests(unittest.TestCase):
    """三级切点检测的方式标记与 cut_detection 对象。"""

    def test_extract_shots_method_no_sampling(self):
        with mock.patch.object(analyze_film, "detect_shots_extract", return_value=[(0.0, 4.0)]):
            segs, method = analyze_film.detect_shots("dummy.mp4", "ffmpeg")
        self.assertEqual(segs, [(0.0, 4.0)])
        self.assertEqual(method, "extract_shots")
        cd = analyze_film.build_cut_detection(method, 13.0, 2.0, 0)
        self.assertNotIn("sampling", cd)     # 正常检测不标采样

    def test_ffmpeg_fallback_method_no_sampling(self):
        with mock.patch.object(analyze_film, "detect_shots_extract", return_value=None), \
             mock.patch.object(analyze_film, "detect_shots_ffmpeg", return_value=[(0.0, 4.0)]):
            segs, method = analyze_film.detect_shots("dummy.mp4", "ffmpeg")
        self.assertEqual(method, "ffmpeg_scene")
        cd = analyze_film.build_cut_detection(method, 13.0, 2.0, 0)
        self.assertNotIn("sampling", cd)

    def test_uniform_fallback_marks_sampling(self):
        with mock.patch.object(analyze_film, "detect_shots_extract", return_value=None), \
             mock.patch.object(analyze_film, "detect_shots_ffmpeg", return_value=None), \
             mock.patch.object(analyze_film, "detect_shots_uniform", return_value=[(0.0, 5.0)]):
            segs, method = analyze_film.detect_shots("dummy.mp4", "ffmpeg")
        self.assertEqual(segs, [(0.0, 5.0)])
        self.assertEqual(method, "uniform5s")
        cd = analyze_film.build_cut_detection(method, 13.0, 2.0, 0)
        self.assertEqual(cd["method"], "uniform5s")
        self.assertIs(cd["sampling"], True)  # 均匀采样分段，非真实镜头边界

    def test_all_detectors_fail_returns_none(self):
        with mock.patch.object(analyze_film, "detect_shots_extract", return_value=None), \
             mock.patch.object(analyze_film, "detect_shots_ffmpeg", return_value=None), \
             mock.patch.object(analyze_film, "detect_shots_uniform", return_value=None):
            segs, method = analyze_film.detect_shots("dummy.mp4", "ffmpeg")
        self.assertIsNone(segs)
        self.assertIsNone(method)

    def test_cut_detection_records_params_and_merged_count(self):
        cd = analyze_film.build_cut_detection("extract_shots", 9.5, 1.5, 3)
        self.assertEqual(cd, {"method": "extract_shots", "thresh": 9.5, "min_dur": 1.5, "merged": 3})


def _cfg(shot_over=None, **top):
    """最小合法 analysis 字典。"""
    shot = {"id": "S1", "t_in": 0.0, "t_out": 3.2, "duration": 3.2,
            "shot_size": "中景", "camera_move": "固定", "angle": "平视",
            "transition": "硬切", "lighting": "", "action": "", "story": "",
            "dialogue": [], "prompt_cn": "", "keyframes": []}
    if shot_over:
        shot.update(shot_over)
    cfg = {"name": "t", "version": 3, "source": "x.mp4",
           "created_at": "2026-01-01T00:00:00", "engine": "none", "shots": [shot]}
    cfg.update(top)
    return cfg


class ValidateAnalysisEvidenceTests(unittest.TestCase):
    """校验器：「不确定」合法、sampling 醒目警告（不阻断）、merged_from 容忍+结构校验。"""

    def test_vocab_sync_between_generator_and_validator(self):
        # 生成端与校验端词表必须一致（防漂移）
        self.assertEqual(analyze_film.VOCAB, validate_analysis.VOCAB)

    def test_buqueding_accepted_without_vocab_warning(self):
        cfg = _cfg({"shot_size": "不确定", "camera_move": "不确定",
                    "angle": "不确定", "transition": "不确定"})
        errors, warnings = validate_analysis.validate(cfg)
        self.assertEqual(errors, [])
        self.assertFalse(any("不在受控词表" in w for w in warnings))

    def test_old_vocab_values_still_accepted(self):
        errors, _ = validate_analysis.validate(_cfg())
        self.assertEqual(errors, [])

    def test_old_analysis_without_new_fields_stays_valid(self):
        # 存量旧 analysis.json：无 cut_detection/merged_from，不报错且无采样警告
        errors, warnings = validate_analysis.validate(_cfg())
        self.assertEqual(errors, [])
        self.assertFalse(any("sampling" in w for w in warnings))

    def test_sampling_warns_but_not_blocks(self):
        cfg = _cfg(cut_detection={"method": "uniform5s", "thresh": 13.0,
                                  "min_dur": 2.0, "merged": 0, "sampling": True})
        errors, warnings = validate_analysis.validate(cfg)
        self.assertEqual(errors, [])
        self.assertTrue(any("sampling" in w and "均匀 5s 采样分段" in w for w in warnings))

    def test_normal_cut_detection_no_sampling_warning(self):
        cfg = _cfg(cut_detection={"method": "extract_shots", "thresh": 13.0,
                                  "min_dur": 2.0, "merged": 2})
        errors, warnings = validate_analysis.validate(cfg)
        self.assertEqual(errors, [])
        self.assertFalse(any("sampling" in w for w in warnings))

    def test_unknown_cut_method_warns_not_blocks(self):
        cfg = _cfg(cut_detection={"method": "magic"})
        errors, warnings = validate_analysis.validate(cfg)
        self.assertEqual(errors, [])
        self.assertTrue(any("method" in w for w in warnings))

    def test_cut_detection_wrong_type_is_error(self):
        errors, _ = validate_analysis.validate(_cfg(cut_detection="uniform5s"))
        self.assertTrue(any("cut_detection" in e for e in errors))

    def test_merged_from_tolerated(self):
        cfg = _cfg({"merged_from": [[0.0, 0.8], [0.8, 1.6], [1.6, 3.2]]})
        errors, _ = validate_analysis.validate(cfg)
        self.assertEqual(errors, [])

    def test_merged_from_malformed_is_error(self):
        errors, _ = validate_analysis.validate(_cfg({"merged_from": "x"}))
        self.assertTrue(any("merged_from" in e for e in errors))
        errors, _ = validate_analysis.validate(_cfg({"merged_from": [["a", 1.0]]}))
        self.assertTrue(any("merged_from" in e for e in errors))
        errors, _ = validate_analysis.validate(_cfg({"merged_from": [[1.0, 1.0]]}))
        self.assertTrue(any("merged_from" in e for e in errors))

    def test_apply_ai_result_accepts_buqueding(self):
        # E03 生成端：AI 填「不确定」直接入库，不再被词表校验顶回
        s = {"id": "S1", "t_in": 0.0, "t_out": 3.0}
        analyze_film.apply_ai_result(s, {"shot_size": "不确定", "camera_move": "不确定",
                                         "angle": "不确定", "transition": "不确定"})
        self.assertEqual(s["shot_size"], "不确定")
        self.assertEqual(s["camera_move"], "不确定")
        self.assertEqual(s["angle"], "不确定")
        self.assertEqual(s["transition"], "不确定")

    def test_boundary_frame_keyframes_tolerated(self):
        # E03 增强：keyframes 数组扩展交界帧命名（S{n}_prev_tail/S{n}_head），校验不报错
        cfg = _cfg({"keyframes": ["keyframes/S2_prev_tail.jpg", "keyframes/S2_head.jpg",
                                  "keyframes/S2_a.jpg", "keyframes/S2_b.jpg",
                                  "keyframes/S2_c.jpg"]})
        errors, _ = validate_analysis.validate(cfg)
        self.assertEqual(errors, [])


class StoryboardEvidenceMappingTests(unittest.TestCase):
    """白模桥接：「不确定」/空值落默认几何（中景/固定/平视），merged_from 透传。"""

    def _convert(self, shot):
        ana = {"shots": [shot]}
        actors, id_of = a2s.build_actors([], "stand")
        return a2s.convert(ana, [], id_of, actors, "stand")

    def _shot(self, **over):
        s = {"id": "S1", "t_in": 0.0, "t_out": 4.0, "duration": 4.0,
             "shot_size": "不确定", "camera_move": "不确定", "angle": "不确定",
             "transition": "不确定", "action": "", "story": "", "prompt_cn": ""}
        s.update(over)
        return s

    def test_buqueding_falls_back_to_default_geometry(self):
        out = self._convert(self._shot())
        self.assertEqual(len(out), 1)
        sh = out[0]
        self.assertEqual(sh["move"], "中景·固定·平视")   # 默认档标签
        self.assertEqual(sh["cam"], "wide")              # 无说话人 -> wide（全景兜底几何）
        self.assertEqual(sh["pos"], [0.0, 2.0, -6.2])    # WIDE_GEO 全景档
        self.assertEqual(sh["look"], [0.0, 1.1, 0.0])
        self.assertEqual(sh["fov"], 46)                  # 中景 fov（stand）
        for k in ("dolly", "pan", "truck", "orbit", "crane", "zoom", "whip"):
            self.assertNotIn(k, sh)                      # 固定档不产生任何运镜参数

    def test_empty_values_fall_back_to_same_defaults(self):
        out = self._convert(self._shot(shot_size="", camera_move="", angle=""))
        self.assertEqual(out[0]["move"], "中景·固定·平视")

    def test_explicit_values_still_map(self):
        # 正常值不受默认值逻辑影响：推->dolly in；仰视->相机降到 0.5m
        out = self._convert(self._shot(shot_size="远景", camera_move="推", angle="仰视"))
        sh = out[0]
        self.assertEqual(sh["dolly"], "in")
        self.assertEqual(sh["pos"][1], 0.5)
        self.assertEqual(sh["move"], "远景·推·仰视")

    def test_merged_from_passthrough_when_multi_segment(self):
        mf = [[0.0, 0.8], [0.8, 1.6], [1.6, 4.0]]
        out = self._convert(self._shot(merged_from=mf))
        self.assertEqual(out[0]["merged_from"], mf)

    def test_merged_from_not_passed_when_single_or_absent(self):
        out = self._convert(self._shot(merged_from=[[0.0, 4.0]]))
        self.assertNotIn("merged_from", out[0])
        out = self._convert(self._shot())
        self.assertNotIn("merged_from", out[0])


class BoundaryFramePlanTests(unittest.TestCase):
    """交界帧计划：命名/位置计算、clamp、首镜无 prev_tail、合并段按段首/段尾取。"""

    def test_names_and_times(self):
        times = [(0.0, 4.0), (4.0, 8.0)]
        plan = analyze_film.boundary_frame_plan(times, 1)
        self.assertEqual(plan, [("keyframes/S2_prev_tail.jpg", 3.85),
                                ("keyframes/S2_head.jpg", 4.05)])

    def test_first_shot_has_no_boundary(self):
        self.assertEqual(analyze_film.boundary_frame_plan([(0.0, 4.0), (4.0, 8.0)], 0), [])

    def test_prev_tail_clamped_to_zero(self):
        # 上一镜只有 0.1s：结尾前 0.15s 越出片头 -> clamp 到 0
        plan = analyze_film.boundary_frame_plan([(0.0, 0.1), (0.1, 5.0)], 1)
        self.assertEqual(plan[0], ("keyframes/S2_prev_tail.jpg", 0.0))
        self.assertEqual(plan[1], ("keyframes/S2_head.jpg", 0.15))

    def test_head_clamped_inside_tiny_shot(self):
        # 本镜只有 0.03s：开头 +0.05s 越过本镜结尾 -> clamp 到 t_out-0.02
        plan = analyze_film.boundary_frame_plan([(0.0, 4.0), (4.0, 4.03)], 1)
        self.assertEqual(plan[1], ("keyframes/S2_head.jpg", 4.01))

    def test_times_clamped_to_video_duration(self):
        plan = analyze_film.boundary_frame_plan([(0.0, 4.0), (4.0, 8.0)], 1, vdur=3.0)
        self.assertEqual(plan[0][1], 3.0)
        self.assertEqual(plan[1][1], 3.0)

    def test_merged_segment_uses_segment_head(self):
        # 合并段：head 取分析段段首（=首个原始段的头），段内切点不用
        segs, mf = analyze_film.merge_short_shots([(0.0, 5.0), (5.0, 5.8), (5.8, 9.0)], 2.0)
        self.assertEqual(segs, [(0.0, 5.8), (5.8, 9.0)])
        plan = analyze_film.boundary_frame_plan(segs, 1)
        self.assertEqual(plan[0][1], 5.65)   # 段首前 0.15s
        self.assertEqual(plan[1][1], 5.85)   # 段首 +0.05s（不是段内任何旧切点）


class SplitBoundaryFramesTests(unittest.TestCase):
    """keyframes 命名分类：交界帧 vs 镜内帧。"""

    def test_split_and_order(self):
        kf = ["keyframes/S2_head.jpg", "keyframes/S2_a.jpg",
              "keyframes/S2_prev_tail.jpg", "keyframes/S2_b.jpg"]
        b, m = analyze_film.split_boundary_frames(kf)
        self.assertEqual(b, ["keyframes/S2_prev_tail.jpg", "keyframes/S2_head.jpg"])
        self.assertEqual(m, ["keyframes/S2_a.jpg", "keyframes/S2_b.jpg"])

    def test_old_three_frame_list_has_no_boundary(self):
        b, m = analyze_film.split_boundary_frames(
            ["keyframes/S1_a.jpg", "keyframes/S1_b.jpg", "keyframes/S1_c.jpg"])
        self.assertEqual(b, [])
        self.assertEqual(len(m), 3)


class BuildAiPromptTests(unittest.TestCase):
    """提示词分支：有交界帧时 transition 要求依证据作答，缺失时允许「不确定」。"""

    def test_full_boundary_requires_evidence(self):
        p = analyze_film.build_ai_prompt(2)
        self.assertIn("前 2 帧是上一镜结尾与本镜开头", p)
        self.assertIn("必须依据这两帧的", p)
        self.assertIn("只有交界帧缺失时才允许", p)
        self.assertIn("硬切", p)          # 词表转场值仍在
        self.assertIn("不确定", p)        # 「不确定」仍是合法值

    def test_partial_boundary_allows_uncertain(self):
        p = analyze_film.build_ai_prompt(1)
        self.assertIn("1 张交界帧", p)
        self.assertIn("证据不充分", p)

    def test_no_boundary_is_module_prompt(self):
        # n<=0 走模块级 AI_PROMPT（可能被 skills/system/fill.md 覆盖层替换）：只断言同一对象
        self.assertIs(analyze_film.build_ai_prompt(0), analyze_film.AI_PROMPT)
        self.assertIs(analyze_film.build_ai_prompt(-1), analyze_film.AI_PROMPT)

    def test_no_boundary_rule_allows_first_shot_none(self):
        # 无交界帧档的转场判据：证据不足填「不确定」，全片第一镜可填「无」
        self.assertIn("一律填 \"不确定\"", analyze_film._TR_RULE_NONE)
        self.assertIn("全片第一镜", analyze_film._TR_RULE_NONE)
        self.assertIn("\"无\"", analyze_film._TR_RULE_NONE)

    def test_vendor_call_uses_boundary_prompt(self):
        # AI 调用接线：n_boundary=2 时请求文本必须是交界帧版提示词
        sent = {}

        class FakeClient:
            def image_part(self, fp):
                return {"type": "image_url", "image_url": fp}

            def chat(self, msgs, **kw):
                sent["text"] = msgs[0]["content"][0]["text"]
                sent["n_images"] = len(msgs[0]["content"]) - 1
                return '{"transition":"硬切"}'

        r = analyze_film.ai_analyze_shot_vendor(
            FakeClient(), ["keyframes/S2_prev_tail.jpg", "keyframes/S2_head.jpg",
                           "keyframes/S2_a.jpg", "keyframes/S2_b.jpg", "keyframes/S2_c.jpg"], 2)
        self.assertEqual(r["transition"], "硬切")
        self.assertIn("前 2 帧是上一镜结尾与本镜开头", sent["text"])
        self.assertEqual(sent["n_images"], 5)


class ExtractKeyframesTests(unittest.TestCase):
    """抽帧落盘：mock _grab_frame 验证命名/顺序/失败容忍（不调真 ffmpeg）。"""

    def test_boundary_frames_included_and_ordered(self):
        with tempfile.TemporaryDirectory() as td:
            calls = []
            with mock.patch.object(analyze_film, "_grab_frame",
                                   side_effect=lambda ff, v, t, p: calls.append((t, p)) or True):
                frames = analyze_film.extract_keyframes("v.mp4", [(0.0, 4.0), (4.0, 8.0)], td, "ff")
        self.assertEqual(frames[0], ["keyframes/S1_a.jpg", "keyframes/S1_b.jpg",
                                     "keyframes/S1_c.jpg"])   # 首镜无交界帧
        self.assertEqual(frames[1], ["keyframes/S2_prev_tail.jpg", "keyframes/S2_head.jpg",
                                     "keyframes/S2_a.jpg", "keyframes/S2_b.jpg",
                                     "keyframes/S2_c.jpg"])   # 交界帧在前
        self.assertEqual(calls[0][0], 1.0)                    # S1 25% 处
        self.assertEqual(calls[3][0], 3.85)                   # S2 prev_tail
        self.assertEqual(calls[4][0], 4.05)                   # S2 head

    def test_transition_frames_disabled(self):
        with tempfile.TemporaryDirectory() as td:
            with mock.patch.object(analyze_film, "_grab_frame", return_value=True):
                frames = analyze_film.extract_keyframes("v.mp4", [(0.0, 4.0), (4.0, 8.0)],
                                                        td, "ff", transition_frames=False)
        self.assertEqual(len(frames[0]), 3)
        self.assertEqual(len(frames[1]), 3)
        self.assertFalse(any("prev_tail" in f0 or "head" in f0 for f0 in frames[1]))

    def test_grab_failure_tolerated(self):
        # 交界帧抽失败：警告跳过，镜内帧照常，不阻断
        def fake(ff, v, t, p):
            return not p.endswith("_head.jpg")
        with tempfile.TemporaryDirectory() as td:
            with mock.patch.object(analyze_film, "_grab_frame", side_effect=fake):
                frames = analyze_film.extract_keyframes("v.mp4", [(0.0, 4.0), (4.0, 8.0)], td, "ff")
        self.assertEqual(frames[1], ["keyframes/S2_prev_tail.jpg",
                                     "keyframes/S2_a.jpg", "keyframes/S2_b.jpg",
                                     "keyframes/S2_c.jpg"])


class EnsureBoundaryFramesTests(unittest.TestCase):
    """fill 模式交界帧补抽与退化路径。"""

    def _setup(self, td, boundary=()):
        ad = Path(td)
        (ad / "keyframes").mkdir()
        shots = [{"id": "S1", "t_in": 0.0, "t_out": 4.0,
                  "keyframes": ["keyframes/S1_a.jpg"]},
                 {"id": "S2", "t_in": 4.0, "t_out": 8.0,
                  "keyframes": ["keyframes/S2_a.jpg"] + list(boundary)}]
        for s in shots:
            for f0 in s["keyframes"]:
                (ad / f0).write_bytes(b"x")
        analysis = {"source": str(ad / "v.mp4"), "shots": shots}
        return ad, analysis, shots

    def test_existing_boundary_reused_without_video(self):
        with tempfile.TemporaryDirectory() as td:
            ad, ana, shots = self._setup(td, ["keyframes/S2_prev_tail.jpg",
                                              "keyframes/S2_head.jpg"])
            ana["source"] = "不存在.mp4"   # 视频不可用也不影响已有交界帧
            out = analyze_film.ensure_boundary_frames(ana, shots, 1, str(ad))
        self.assertEqual(out, ["keyframes/S2_prev_tail.jpg", "keyframes/S2_head.jpg"])

    def test_first_shot_returns_empty(self):
        with tempfile.TemporaryDirectory() as td:
            ad, ana, shots = self._setup(td)
            out = analyze_film.ensure_boundary_frames(ana, shots, 0, str(ad))
        self.assertEqual(out, [])

    def test_missing_video_degrades_to_empty(self):
        with tempfile.TemporaryDirectory() as td:
            ad, ana, shots = self._setup(td)
            ana["source"] = "不存在.mp4"
            out = analyze_film.ensure_boundary_frames(ana, shots, 1, str(ad))
        self.assertEqual(out, [])                                # 退化：提示词允许「不确定」
        self.assertNotIn("keyframes/S2_head.jpg", shots[1]["keyframes"])   # 不写假数据

    def test_missing_video_keeps_existing_partial(self):
        with tempfile.TemporaryDirectory() as td:
            ad, ana, shots = self._setup(td, ["keyframes/S2_prev_tail.jpg"])
            ana["source"] = "不存在.mp4"
            out = analyze_film.ensure_boundary_frames(ana, shots, 1, str(ad))
        self.assertEqual(out, ["keyframes/S2_prev_tail.jpg"])    # 已有 1 张照用

    def test_extracts_missing_and_writes_back(self):
        with tempfile.TemporaryDirectory() as td:
            ad, ana, shots = self._setup(td)
            (ad / "v.mp4").write_bytes(b"x")                     # 源视频可用
            grabbed = []
            def fake(ff, v, t, p):
                grabbed.append((t, p)); return True
            with mock.patch.object(analyze_film, "_grab_frame", side_effect=fake):
                out = analyze_film.ensure_boundary_frames(ana, shots, 1, str(ad), ff="ff")
        self.assertEqual(out, ["keyframes/S2_prev_tail.jpg", "keyframes/S2_head.jpg"])
        self.assertEqual([t for t, _ in grabbed], [3.85, 4.05])  # 按 shots 时间轴计算
        self.assertIn("keyframes/S2_prev_tail.jpg", shots[1]["keyframes"])  # 写回持久化
        self.assertIn("keyframes/S2_head.jpg", shots[1]["keyframes"])

    def test_partial_completion_keeps_order(self):
        # 已有 head、缺 prev_tail：补抽后固定 prev_tail 在前
        with tempfile.TemporaryDirectory() as td:
            ad, ana, shots = self._setup(td, ["keyframes/S2_head.jpg"])
            (ad / "v.mp4").write_bytes(b"x")
            with mock.patch.object(analyze_film, "_grab_frame", return_value=True):
                out = analyze_film.ensure_boundary_frames(ana, shots, 1, str(ad), ff="ff")
        self.assertEqual(out, ["keyframes/S2_prev_tail.jpg", "keyframes/S2_head.jpg"])


class StoryboardOriginTests(unittest.TestCase):
    """E09：题材模板假设显式标记——镜级 origin.fields + 顶层 _generator + schematic/faithful 双模式。"""

    LINES2 = [{"t_in": 0.2, "t_out": 2.0, "text": "曹操台词", "speaker": "曹操"},
              {"t_in": 5.0, "t_out": 7.0, "text": "王朗台词", "speaker": "王朗"}]

    def _shot(self, **over):
        s = {"id": "S1", "t_in": 0.0, "t_out": 4.0, "duration": 4.0,
             "shot_size": "特写", "camera_move": "固定", "angle": "平视",
             "transition": "硬切", "action": "", "story": "室内议事", "prompt_cn": ""}
        s.update(over)
        return s

    def _convert(self, shots, mode="schematic", lines=None, assign_pos=True):
        ana = {"shots": shots}
        lines = self.LINES2 if lines is None else lines
        actors, id_of = a2s.build_actors(lines, "stand", assign_pos=assign_pos)
        out = a2s.convert(ana, lines, id_of, actors, "stand", mode=mode)
        return out, actors

    # ---- schematic（默认）：现状行为 + 来源标记 ----

    def test_schematic_origin_covers_scene_staging_pose_camera(self):
        out, _ = self._convert([self._shot()])
        sh = out[0]
        f = sh["origin"]["fields"]
        for k in ("scene", "staging", "pose", "camera", "on_set", "lines"):
            self.assertIn(k, f)                       # 字段覆盖
        self.assertEqual(f["scene"], "proposal")      # 关键词推断=提案
        self.assertEqual(f["staging"], "proposal")    # 画外移出=提案
        self.assertEqual(f["pose"], "proposal")
        self.assertEqual(f["camera"], "analysis_mapped")   # 识别值经规则映射
        self.assertEqual(f["lines"], "evidence")      # 台词来自台词脚本
        self.assertIn("非原片", sh["origin"]["note"])
        # schematic 行为不变：staging/scene/pose/pos 照常写出
        self.assertIn("staging", sh)
        self.assertEqual(sh["scene"], "room")
        self.assertIn("pose", sh)
        self.assertIn("pos", sh)

    def test_schematic_camera_default_when_all_unknown(self):
        out, _ = self._convert([self._shot(shot_size="不确定", camera_move="", angle="不确定")])
        self.assertEqual(out[0]["origin"]["fields"]["camera"], "default")

    def test_scene_evidence_from_explicit_analysis_field(self):
        out, _ = self._convert([self._shot(scene="field")])
        sh = out[0]
        self.assertEqual(sh["scene"], "field")                    # 采用证据值而非关键词推断
        self.assertEqual(sh["origin"]["fields"]["scene"], "evidence")

    def test_generator_block_schematic(self):
        out, _ = self._convert([self._shot()])
        blk = a2s.generator_block("schematic", out)
        self.assertEqual(blk["tool"], "analysis_to_storyboard")
        self.assertEqual(blk["mode"], "示意预演")
        self.assertEqual(blk["mode_code"], "schematic")
        self.assertIn("非原片复原", blk["disclaimer"])
        self.assertNotIn("needs_review", blk)                     # schematic 无待确认项

    # ---- faithful：只用有证据字段，提案剔除留空 ----

    def test_faithful_actors_have_no_pos(self):
        _, actors = self._convert([self._shot()], mode="faithful", assign_pos=False)
        for ac in actors.values():
            self.assertNotIn("pos", ac)                           # 站位不自动排

    def test_faithful_cu_omits_geometry_and_marks_needs(self):
        out, _ = self._convert([self._shot()], mode="faithful", assign_pos=False)
        sh = out[0]
        self.assertEqual(sh["cam"], "cu")                         # cam 语义保留（台词=证据）
        for k in ("pos", "look", "fov", "scene", "pose", "staging"):
            self.assertNotIn(k, sh)                               # 提案/依赖站位的几何全留空
        f = sh["origin"]["fields"]
        self.assertEqual(f["camera"], "omitted")
        self.assertEqual(f["scene"], "omitted")
        self.assertEqual(f["pose"], "omitted")
        self.assertEqual(f["staging"], "omitted")
        self.assertEqual(f["lines"], "evidence")
        self.assertEqual(sorted(sh["origin"]["needs"]), ["camera", "scene", "staging"])

    def test_faithful_wide_with_real_size_maps_camera(self):
        # wide 不依赖站位且景别有证据 -> 几何照出，标 analysis_mapped
        shot = self._shot(shot_size="远景", t_in=10.0, t_out=14.0, duration=4.0)
        out, _ = self._convert([shot], mode="faithful", assign_pos=False)
        sh = out[0]
        self.assertEqual(sh["cam"], "wide")
        self.assertIn("pos", sh)
        self.assertEqual(sh["origin"]["fields"]["camera"], "analysis_mapped")
        self.assertNotIn("camera", sh["origin"].get("needs", []))

    def test_faithful_wide_without_size_evidence_omits_camera(self):
        shot = self._shot(shot_size="不确定", camera_move="不确定", angle="不确定",
                          t_in=10.0, t_out=14.0, duration=4.0)
        out, _ = self._convert([shot], mode="faithful", assign_pos=False)
        sh = out[0]
        self.assertNotIn("pos", sh)
        self.assertEqual(sh["origin"]["fields"]["camera"], "omitted")
        self.assertIn("camera", sh["origin"]["needs"])

    def test_faithful_scene_evidence_kept(self):
        out, _ = self._convert([self._shot(scene="room")], mode="faithful", assign_pos=False)
        sh = out[0]
        self.assertEqual(sh["scene"], "room")
        self.assertEqual(sh["origin"]["fields"]["scene"], "evidence")
        self.assertNotIn("scene", sh["origin"].get("needs", []))

    def test_generator_block_faithful_aggregates_needs(self):
        out, _ = self._convert([self._shot()], mode="faithful", assign_pos=False)
        blk = a2s.generator_block("faithful", out)
        self.assertEqual(blk["mode"], "忠实重建")
        self.assertIn("人工确认", blk["disclaimer"])
        self.assertEqual(blk["needs_review"], {"S1": ["staging", "camera", "scene"]})

    # ---- 校验器容忍：origin/_generator 未知字段不报错 ----

    def test_validate_dialogue_tolerates_origin_and_generator(self):
        import validate_dialogue
        board = {"actors": {"c": {"name": "甲", "shirt": [1, 2, 3], "pos": [0, 0], "static": True}},
                 "shots": [{"id": "S1", "dur": 3.0, "cam": "wide",
                            "pos": [0, 2, -6], "look": [0, 1, 0], "fov": 48,
                            "origin": {"fields": {"camera": "default"}, "note": "测试"}}],
                 "_generator": {"tool": "analysis_to_storyboard", "mode": "示意预演",
                                "disclaimer": "几何/站位为程序提案，非原片复原"}}
        r = validate_dialogue.validate_document(board)
        self.assertEqual(r["errors"], [])


class DialogueTrackTests(unittest.TestCase):
    """E08 台词独立音轨：顶层 dialogue_track=完整时间本体（一句可跨多镜），
    shot.dialogue=交集引用视图（span=primary|overlap，显示时间裁镜内、本体不裁）。"""

    def _project(self, td, lines):
        """建临时项目：拉片/v1/analysis.json（0-4/4-8/8-12 三镜）+ 台词/台词脚本.json。"""
        proj = Path(td)
        adir = proj / "拉片" / "v1"
        adir.mkdir(parents=True)
        (proj / "台词").mkdir()
        ana = {"name": "v1", "version": 3, "source": "x.mp4",
               "created_at": "2026-01-01T00:00:00", "engine": "none",
               "shots": [{"id": f"S{i+1}", "t_in": i * 4.0, "t_out": (i + 1) * 4.0,
                          "duration": 4.0, "shot_size": "中景", "camera_move": "固定",
                          "angle": "平视", "transition": "硬切", "dialogue": [],
                          "keyframes": []} for i in range(3)]}
        (adir / "analysis.json").write_text(json.dumps(ana, ensure_ascii=False), encoding="utf-8")
        (proj / "台词" / "台词脚本.json").write_text(
            json.dumps({"speakers": {}, "lines": lines}, ensure_ascii=False), encoding="utf-8")
        return proj, ana

    def test_cross_shot_line_spans_all_overlapped_shots(self):
        """跨镜长句 2-10 挂三镜：S1 overlap(2-4) / S2 primary(4-8，中点6) / S3 overlap(8-10)；
        本体完整不裁；引用视图通过校验。"""
        lines = [{"t_in": 2.0, "t_out": 10.0, "speaker": "甲", "text": "跨镜长句"}]
        with tempfile.TemporaryDirectory() as td:
            proj, ana = self._project(td, lines)
            n = analyze_film.merge_script_lines(ana, str(proj))
        self.assertEqual(n, 1)
        self.assertEqual(ana["dialogue_track"],
                         [{"event": "L001", "t_in": 2.0, "t_out": 10.0,
                           "speaker": "甲", "text": "跨镜长句"}])      # 本体不裁
        s1, s2, s3 = ana["shots"]
        self.assertEqual([(d["t_in"], d["t_out"], d["span"], d["event"]) for d in s1["dialogue"]],
                         [(2.0, 4.0, "overlap", "L001")])
        self.assertEqual([(d["t_in"], d["t_out"], d["span"]) for d in s2["dialogue"]],
                         [(4.0, 8.0, "primary")])                      # 中点 6.0 在 S2
        self.assertEqual([(d["t_in"], d["t_out"], d["span"]) for d in s3["dialogue"]],
                         [(8.0, 10.0, "overlap")])
        errors, _ = validate_analysis.validate(ana)
        self.assertEqual(errors, [])

    def test_midpoint_in_gap_falls_back_to_max_overlap(self):
        """中点落镜缝：主场取重叠最大镜，validate 不报错（缝隙 fallback 合法）。"""
        lines = [{"t_in": 2.5, "t_out": 4.5, "speaker": "甲", "text": "缝隙句"}]
        with tempfile.TemporaryDirectory() as td:
            proj, ana = self._project(td, lines)
            ana["shots"] = ana["shots"][:2]                            # 0-4 / 4-8 两镜
            ana["shots"][0]["t_out"] = 3.0; ana["shots"][0]["duration"] = 3.0   # 造缝 3-4
            analyze_film.merge_script_lines(ana, str(proj))
            spans = {s["id"]: [(d["span"], d["event"]) for d in s["dialogue"]] for s in ana["shots"]}
            # 中点 3.5 在缝（3-4），两侧重叠各 0.5 —— max 取先命中者 S1
            self.assertEqual(spans["S1"], [("primary", "L001")])
            self.assertEqual(spans["S2"], [("overlap", "L001")])
            errors, _ = validate_analysis.validate(ana)
        self.assertEqual(errors, [])

    def test_orphan_line_stays_in_track_only(self):
        """孤儿句（与任何镜无交集）：本体保留，不进任何镜的视图。"""
        lines = [{"t_in": 50.0, "t_out": 52.0, "speaker": "甲", "text": "片外句"}]
        with tempfile.TemporaryDirectory() as td:
            proj, ana = self._project(td, lines)
            n = analyze_film.merge_script_lines(ana, str(proj))
        self.assertEqual(n, 1)
        self.assertEqual(len(ana["dialogue_track"]), 1)
        self.assertTrue(all(s["dialogue"] == [] for s in ana["shots"]))

    def test_resync_event_ids_stable_and_text_change_gets_new_id(self):
        """幂等：重同步 id 不变；改一条文本 → 该条换新 id（递增不复用）、其余保留。"""
        lines = [{"t_in": 0.5, "t_out": 2.0, "speaker": "甲", "text": "第一句"},
                 {"t_in": 5.0, "t_out": 7.0, "speaker": "乙", "text": "第二句"}]
        with tempfile.TemporaryDirectory() as td:
            proj, ana = self._project(td, lines)
            analyze_film.merge_script_lines(ana, str(proj))
            ids1 = [e["event"] for e in ana["dialogue_track"]]
            analyze_film.merge_script_lines(ana, str(proj))            # 重同步
            ids2 = [e["event"] for e in ana["dialogue_track"]]
            self.assertEqual(ids1, ids2)
            # 改第一条文本后重同步
            sc_p = proj / "台词" / "台词脚本.json"
            sc = json.loads(sc_p.read_text(encoding="utf-8"))
            sc["lines"][0]["text"] = "第一句改"
            sc_p.write_text(json.dumps(sc, ensure_ascii=False), encoding="utf-8")
            analyze_film.merge_script_lines(ana, str(proj))
            ids3 = [e["event"] for e in ana["dialogue_track"]]
        self.assertEqual(ids3[1], ids1[1])                             # 未变的保留
        self.assertNotEqual(ids3[0], ids1[0])                          # 改过的换新
        self.assertEqual(ids3[0], "L003")                              # 新 id 递增不复用

    # ---- validate_analysis：dialogue_track 结构 + span 视图 + 旧格式兼容 ----

    def test_validate_dialogue_track_structure(self):
        tr = lambda **kw: dict({"event": "L001", "t_in": 1.0, "t_out": 2.0,
                                "speaker": "甲", "text": "x"}, **kw)
        errors, _ = validate_analysis.validate(_cfg(dialogue_track=[tr()]))
        self.assertEqual(errors, [])                                   # 合法
        errors, _ = validate_analysis.validate(_cfg(dialogue_track=[tr(), tr()]))
        self.assertTrue(any("event 重复" in e for e in errors))
        errors, _ = validate_analysis.validate(
            _cfg(dialogue_track=[tr(event="L001", t_in=2.0, t_out=3.0),
                                 tr(event="L002", t_in=1.0, t_out=1.5)]))
        self.assertTrue(any("须按时间排序" in e for e in errors))
        errors, _ = validate_analysis.validate(_cfg(dialogue_track=["x"]))
        self.assertTrue(any("应为对象" in e for e in errors))
        errors, _ = validate_analysis.validate(_cfg(dialogue_track=[tr(t_in=2.0, t_out=2.0)]))
        self.assertTrue(any("须满足 t_in < t_out" in e for e in errors))

    def test_validate_span_rules(self):
        # span 非法值
        d = {"speaker": "甲", "text": "x", "t_in": 0.5, "t_out": 2.0, "span": "bad"}
        errors, _ = validate_analysis.validate(_cfg({"dialogue": [d]}))
        self.assertTrue(any("span 应为 primary|overlap" in e for e in errors))
        # span 项与镜头区间无交集（镜头 0-3.2）
        d = {"speaker": "甲", "text": "x", "t_in": 5.0, "t_out": 6.0, "span": "overlap"}
        errors, _ = validate_analysis.validate(_cfg({"dialogue": [d]}))
        self.assertTrue(any("无交集" in e for e in errors))
        # span 项有交集即过（显示时间允许裁在镜内，不要求完整落镜内）
        d = {"speaker": "甲", "text": "x", "t_in": 3.0, "t_out": 3.2, "span": "overlap"}
        errors, _ = validate_analysis.validate(_cfg({"dialogue": [d]}))
        self.assertEqual(errors, [])

    def _two_shot_cfg(self, d1=None, d2=None, track=None):
        """两镜 0-3 / 4-7（有缝 3-4）的最小合法 cfg。"""
        cfg = _cfg()
        s1 = dict(cfg["shots"][0], id="S1", t_in=0.0, t_out=3.0, duration=3.0, dialogue=d1 or [])
        s2 = dict(cfg["shots"][0], id="S2", t_in=4.0, t_out=7.0, duration=3.0, dialogue=d2 or [])
        cfg["shots"] = [s1, s2]
        if track is not None:
            cfg["dialogue_track"] = track
        return cfg

    def test_validate_primary_midpoint_and_dangling_event(self):
        track = [{"event": "L001", "t_in": 2.0, "t_out": 8.0, "speaker": "甲", "text": "x"}]  # 中点 5.0 在 S2
        d_s1 = {"speaker": "甲", "text": "x", "t_in": 2.0, "t_out": 3.0, "span": "primary", "event": "L001"}
        errors, _ = validate_analysis.validate(self._two_shot_cfg(d1=[d_s1], track=track))
        self.assertTrue(any("中点" in e for e in errors))              # primary 挂错镜
        d_s2 = dict(d_s1, t_in=4.0, t_out=7.0)
        errors, _ = validate_analysis.validate(self._two_shot_cfg(d2=[d_s2], track=track))
        self.assertEqual(errors, [])                                   # primary 挂对镜
        d_bad = dict(d_s2, event="L999")
        errors, _ = validate_analysis.validate(self._two_shot_cfg(d2=[d_bad], track=track))
        self.assertTrue(any("不存在" in e for e in errors))            # dangling event

    def test_validate_legacy_dialogue_without_span_still_enclosed(self):
        # 旧格式（无 span）：维持落镜内规则——出界报错、落内通过
        d_in = {"speaker": "甲", "text": "x", "t_in": 0.5, "t_out": 2.0}
        errors, _ = validate_analysis.validate(_cfg({"dialogue": [d_in]}))
        self.assertEqual(errors, [])
        d_out = {"speaker": "甲", "text": "x", "t_in": 0.5, "t_out": 9.0}
        errors, _ = validate_analysis.validate(_cfg({"dialogue": [d_out]}))
        self.assertTrue(any("超出镜头区间" in e for e in errors))

    # ---- 桥接：analysis_to_storyboard 按 dialogue_track 分配 ----

    def _board_shots(self):
        return [{"id": f"S{i+1}", "t_in": i * 4.0, "t_out": (i + 1) * 4.0, "duration": 4.0,
                 "shot_size": "中景", "camera_move": "固定", "angle": "平视",
                 "transition": "硬切", "action": "", "story": "室内议事", "prompt_cn": ""}
                for i in range(3)]

    def test_storyboard_track_line_appears_in_each_spanned_shot(self):
        track = [{"event": "L001", "t_in": 2.0, "t_out": 10.0, "speaker": "曹操", "text": "跨镜"}]
        lines = [{"t_in": 2.0, "t_out": 10.0, "speaker": "曹操", "text": "跨镜"}]
        actors, id_of = a2s.build_actors(lines, "stand")
        out = a2s.convert({"shots": self._board_shots(), "dialogue_track": track},
                          lines, id_of, actors, "stand")
        for i, span in ((0, "overlap"), (1, "primary"), (2, "overlap")):
            ls = out[i].get("lines", [])
            self.assertEqual(len(ls), 1, f"S{i+1} 应有 1 条台词")
            self.assertEqual((ls[0]["event"], ls[0]["span"], ls[0]["line"]),
                             ("L001", span, "跨镜"))
        # origin.fields.lines 维持 evidence
        self.assertEqual(out[1]["origin"]["fields"]["lines"], "evidence")

    def test_speakers_beyond_actor_id_pool_keep_ids(self):
        """超过预设 id 池（10 位）的说话人不得被丢掉：曾截断 → 第 11+ 位台词 speaker=None。"""
        lines = [{"t_in": float(i), "t_out": float(i) + 1.0, "speaker": f"角色{i+1}", "text": f"词{i+1}"}
                 for i in range(13)]
        actors, id_of = a2s.build_actors(lines, "stand")
        self.assertEqual(len(actors), 13)
        self.assertEqual({L["speaker"] for L in lines} - set(id_of), set(), "每位说话人都要有 id")
        self.assertEqual(len(set(actors)), len(actors), "actor id 不得互相冲突")
        self.assertEqual(list(actors)[:len(a2s.ACTOR_IDS)], list(a2s.ACTOR_IDS),
                         "前 10 位仍沿用预设 id 池（不改变既有产物口径）")
        out = a2s.convert({"shots": self._board_shots()}, lines, id_of, actors, "stand")
        speakers = [L.get("speaker") for sh in out for L in sh.get("lines", [])]
        self.assertTrue(speakers, "应有台词落进分镜")
        self.assertNotIn(None, speakers, "speaker 为 None 会让 ④ 音色与 ⑤ 演员层永久认不到")

    def test_storyboard_without_track_keeps_single_assignment(self):
        lines = [{"t_in": 2.0, "t_out": 10.0, "speaker": "曹操", "text": "跨镜"}]
        actors, id_of = a2s.build_actors(lines, "stand")
        out = a2s.convert({"shots": self._board_shots()}, lines, id_of, actors, "stand")
        hits = [(i, L) for i, sh in enumerate(out) for L in sh.get("lines", []) if L["line"] == "跨镜"]
        self.assertEqual(len(hits), 1)                                 # 旧逻辑：只归一镜
        self.assertNotIn("span", hits[0][1])                           # 无 span 字段
        self.assertNotIn("event", hits[0][1])

    # ---- Excel：sheet2 按本体去重、sheet1 跨镜标记 ----

    def test_xlsx_track_sheet2_dedup_and_arrow_marker(self):
        import subprocess
        from openpyxl import load_workbook
        track = [{"event": "L001", "t_in": 2.0, "t_out": 10.0, "speaker": "甲", "text": "跨镜长句"}]
        shots = []
        for i, (span, lo, hi) in enumerate((("overlap", 2.0, 4.0), ("primary", 4.0, 8.0),
                                            ("overlap", 8.0, 10.0))):
            shots.append({"id": f"S{i+1}", "t_in": i * 4.0, "t_out": (i + 1) * 4.0, "duration": 4.0,
                          "shot_size": "中景", "camera_move": "固定", "angle": "平视",
                          "dialogue": [{"event": "L001", "speaker": "甲", "text": "跨镜长句",
                                        "t_in": lo, "t_out": hi, "span": span}]})
        ana = {"name": "t", "version": 3, "source": "x.mp4", "dialogue_track": track, "shots": shots}
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "analysis.json"
            p.write_text(json.dumps(ana, ensure_ascii=False), encoding="utf-8")
            r = subprocess.run(
                [sys.executable, str(ROOT / "workbench" / "tools" / "export_analysis_xlsx.py"), str(p)],
                capture_output=True, text=True, encoding="utf-8", errors="replace")
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            wb = load_workbook(str(Path(td) / "t_分镜脚本.xlsx"))
            ws1, ws2 = wb["分镜脚本"], wb["台词汇总"]
            self.assertEqual(ws2.max_row, 1 + len(track))              # 本体一句一行，天然去重
            self.assertEqual(ws2.cell(2, 1).value, "00:02.0–00:10.0")  # 完整时间
            c1, c2 = ws1.cell(2, 10).value, ws1.cell(3, 10).value      # 台词列
            self.assertTrue(c1.startswith("↔"))                        # overlap 镜带跨镜标记
            self.assertFalse(c2.startswith("↔"))                       # primary 镜不带


class CliSmokeTests(unittest.TestCase):
    """命令行级冒烟：validate_analysis.py 对 sampling 分析退出码 0 且打印醒目警告。"""

    def test_validate_cli_accepts_sampling_analysis(self):
        import subprocess
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "analysis.json"
            p.write_text(json.dumps(
                _cfg(cut_detection={"method": "uniform5s", "thresh": 13.0,
                                    "min_dur": 2.0, "merged": 0, "sampling": True}),
                ensure_ascii=False), encoding="utf-8")
            r = subprocess.run(
                [sys.executable, str(ROOT / "workbench" / "tools" / "validate_analysis.py"), str(p)],
                capture_output=True, text=True, encoding="utf-8", errors="replace")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("均匀 5s 采样分段", r.stdout)


if __name__ == "__main__":
    unittest.main()
